"""Tests for the AWS inventory utility using botocore Stubber (no real AWS calls)."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import boto3
import pytest
from botocore.awsrequest import AWSResponse
from botocore.exceptions import ClientError
from botocore.stub import Stubber
from openpyxl import load_workbook

from src.aws_inventory import session as session_module
from src.aws_inventory.cli import account_id_type, parse_args
from src.aws_inventory.cloudwatch import MetricQuery, latest_activity
from src.aws_inventory.collectors import COLLECTORS
from src.aws_inventory.collectors.messaging import collect_sqs
from src.aws_inventory.exporters import write_json, write_xlsx
from src.aws_inventory.model import Dependency, Inventory, Resource, parse_arn
from src.aws_inventory.runner import _run_collector, merge
from src.aws_inventory.session import AccountMismatchError, Context, build_context

ACCOUNT = "123456789012"
REGION = "us-east-1"


def make_client(service: str) -> Any:
    return boto3.client(
        service, region_name=REGION, aws_access_key_id="test", aws_secret_access_key="test"
    )


class StubbedContext(Context):
    """Context whose clients are pre-built stubbed clients keyed by service name."""

    def __init__(self, clients: dict[str, Any]) -> None:
        super().__init__(boto3.Session(), ACCOUNT)
        self._stubbed = clients

    def client(self, service: str, region: str | None = None) -> Any:
        return self._stubbed[service]


def test_account_id_must_be_twelve_digits() -> None:
    assert account_id_type(ACCOUNT) == ACCOUNT
    for bad in ("123", "12345678901a", "1234567890123"):  # pragma: allowlist secret
        with pytest.raises(Exception):
            account_id_type(bad)
    assert parse_args([ACCOUNT, "--regions", "us-east-1, eu-west-1"]).regions == [
        "us-east-1",
        "eu-west-1",
    ]


def test_all_expected_collectors_registered() -> None:
    assert set(COLLECTORS) == {
        "apigateway",
        "cloudfront",
        "dynamodb",
        "ec2",
        "ecs",
        "eks",
        "elbv2",
        "iam",
        "lambda",
        "rds",
        "s3",
        "sns",
        "sqs",
    }
    assert {name for name, c in COLLECTORS.items() if c.scope == "global"} == {
        "iam",
        "s3",
        "cloudfront",
    }


def test_account_mismatch_aborts(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakeSts:
        def get_caller_identity(self) -> dict[str, str]:
            return {"Account": "999999999999", "Arn": "arn:aws:iam::999999999999:user/someone"}

    class FakeSession:
        def __init__(self, **_: Any) -> None:
            pass

        def client(self, name: str, **_: Any) -> FakeSts:
            return FakeSts()

    monkeypatch.setattr(session_module.boto3, "Session", FakeSession)
    with pytest.raises(AccountMismatchError, match="999999999999"):
        build_context(ACCOUNT)


def test_matching_account_builds_context(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakeSts:
        def get_caller_identity(self) -> dict[str, str]:
            return {"Account": ACCOUNT, "Arn": f"arn:aws-cn:iam::{ACCOUNT}:user/me"}

    class FakeSession:
        def __init__(self, **_: Any) -> None:
            pass

        def client(self, name: str, **_: Any) -> FakeSts:
            return FakeSts()

    monkeypatch.setattr(session_module.boto3, "Session", FakeSession)
    ctx = build_context(ACCOUNT)
    assert ctx.partition == "aws-cn"
    assert ctx.global_region == "cn-north-1"


def test_collector_access_denied_is_recorded_not_raised() -> None:
    def boom(ctx: Context, region: str) -> list[Resource]:
        raise ClientError({"Error": {"Code": "AccessDenied", "Message": "nope"}}, "ListThings")

    from src.aws_inventory.collectors.base import Collector

    ctx = Context(boto3.Session(), ACCOUNT)
    assert _run_collector(ctx, Collector("demo", "regional", boom), REGION) == []
    assert len(ctx.errors) == 1
    assert ctx.errors[0].service == "demo"
    assert "AccessDenied" in ctx.errors[0].message


def test_safe_ignores_expected_codes() -> None:
    ctx = Context(boto3.Session(), ACCOUNT)

    def missing() -> None:
        raise ClientError({"Error": {"Code": "NoSuchTagSet", "Message": ""}}, "GetBucketTagging")

    assert (
        ctx.safe("s3", REGION, "GetBucketTagging", missing, {}, ignore_codes=("NoSuchTagSet",))
        == {}
    )
    assert ctx.errors == []


def test_latest_activity_distinguishes_none_from_unknown() -> None:
    cloudwatch = make_client("cloudwatch")
    ctx = StubbedContext({"cloudwatch": cloudwatch})
    newest = datetime(2026, 10, 20, tzinfo=timezone.utc)
    with Stubber(cloudwatch) as stubber:
        stubber.add_response(
            "get_metric_data",
            {
                "MetricDataResults": [
                    {
                        "Id": "q0",
                        "Timestamps": [newest, datetime(2026, 10, 19, tzinfo=timezone.utc)],
                        "Values": [5.0, 1.0],
                    },
                    {"Id": "q1", "Timestamps": [newest], "Values": [0.0]},
                ]
            },
        )
        result = latest_activity(
            ctx,
            "lambda",
            REGION,
            {
                "busy": MetricQuery("AWS/Lambda", "Invocations", {"FunctionName": "busy"}),
                "idle": MetricQuery("AWS/Lambda", "Invocations", {"FunctionName": "idle"}),
            },
        )
    assert result == {"busy": newest.isoformat(), "idle": None}


def test_latest_activity_failure_leaves_keys_unknown() -> None:
    cloudwatch = make_client("cloudwatch")
    ctx = StubbedContext({"cloudwatch": cloudwatch})
    with Stubber(cloudwatch) as stubber:
        stubber.add_client_error("get_metric_data", service_error_code="AccessDenied")
        result = latest_activity(
            ctx,
            "sqs",
            REGION,
            {"q": MetricQuery("AWS/SQS", "NumberOfMessagesSent", {"QueueName": "q"})},
        )
    assert result == {}
    assert ctx.errors and ctx.errors[0].operation == "cloudwatch:GetMetricData"


def test_sqs_collector_maps_resource_and_dependencies() -> None:
    sqs, cloudwatch = make_client("sqs"), make_client("cloudwatch")
    ctx = StubbedContext({"sqs": sqs, "cloudwatch": cloudwatch})
    url = f"https://sqs.{REGION}.amazonaws.com/{ACCOUNT}/orders"
    dlq = f"arn:aws:sqs:{REGION}:{ACCOUNT}:orders-dlq"
    sent = datetime(2026, 10, 25, tzinfo=timezone.utc)
    with Stubber(sqs) as sqs_stub, Stubber(cloudwatch) as cw_stub:
        sqs_stub.add_response("list_queues", {"QueueUrls": [url]})
        sqs_stub.add_response(
            "get_queue_attributes",
            {
                "Attributes": {
                    "QueueArn": f"arn:aws:sqs:{REGION}:{ACCOUNT}:orders",
                    "CreatedTimestamp": "1700000000",
                    "RedrivePolicy": f'{{"deadLetterTargetArn":"{dlq}","maxReceiveCount":3}}',
                    "ApproximateNumberOfMessages": "4",
                }
            },
        )
        sqs_stub.add_response("list_queue_tags", {"Tags": {"team": "billing"}})
        cw_stub.add_response(
            "get_metric_data",
            {"MetricDataResults": [{"Id": "q0", "Timestamps": [sent], "Values": [12.0]}]},
        )
        resources = collect_sqs(ctx, REGION)

    assert len(resources) == 1
    queue = resources[0]
    assert (queue.service, queue.type, queue.id) == ("sqs", "queue", "orders")
    assert queue.tags == {"team": "billing"}
    assert queue.is_active is True
    assert queue.last_activity == sent.isoformat()
    assert queue.last_activity_source == "cloudwatch:NumberOfMessagesSent"
    assert queue.dependencies == [Dependency("dead-letter-queue", dlq, "sqs")]
    assert ctx.errors == []


def test_parse_arn_and_merge_prefers_deep_results() -> None:
    assert parse_arn("arn:aws:elasticloadbalancing:us-east-1:1:loadbalancer/app/x/1") == (
        "elbv2",
        "us-east-1",
        "loadbalancer",
        "app/x/1",
    )
    assert parse_arn("arn:aws:sqs:us-east-1:1:orders") == ("sqs", "us-east-1", "resource", "orders")

    arn = "arn:aws:sqs:us-east-1:1:orders"
    deep = [Resource("sqs", "queue", "orders", arn, REGION, status="available")]
    stubs = [
        Resource("sqs", "resource", "orders", arn, REGION, status="unknown"),
        Resource(
            "kinesis",
            "stream",
            "s",
            "arn:aws:kinesis:us-east-1:1:stream/s",
            REGION,
            status="unknown",
        ),
    ]
    merged = merge(deep, stubs)
    assert [(r.service, r.status) for r in merged] == [("sqs", "available"), ("kinesis", "unknown")]


def sample_inventory() -> Inventory:
    return Inventory(
        account={"id": ACCOUNT, "regions_scanned": [REGION]},
        resources=[
            Resource(
                "ec2",
                "instance",
                "i-1",
                f"arn:aws:ec2:{REGION}:{ACCOUNT}:instance/i-1",
                REGION,
                name="web",
                status="running",
                is_active=True,
                last_activity="2026-10-27T00:00:00+00:00",
                tags={"env": "prod"},
                metadata={"instance_type": "t3.micro"},
                dependencies=[Dependency("in-vpc", "vpc-1", "ec2")],
            ),
            Resource("s3", "bucket", "b", "arn:aws:s3:::b", REGION, status="available"),
        ],
        cost={
            "by_service": {"ec2": {"total": 12.5, "by_region": {REGION: 12.5}}},
            "unmapped": {"Amazon Kinesis": 3.0},
        },
        services_scanned=["ec2", "s3"],
    )


def test_json_and_excel_export(tmp_path: Path) -> None:
    inventory = sample_inventory()
    json_path = write_json(inventory, tmp_path / "out.json")
    xlsx_path = write_xlsx(inventory, tmp_path / "out.xlsx")

    import json

    data = json.loads(json_path.read_text(encoding="utf-8"))
    assert set(data) >= {"account", "notes", "services", "dependencies", "cost", "coverage"}
    ec2 = data["services"]["ec2"]
    assert ec2["summary"]["resource_count"] == 1
    assert ec2["summary"]["active_count"] == 1
    assert ec2["summary"]["cost_90d"] == 12.5
    assert data["services"]["s3"]["summary"]["unknown_activity_count"] == 1
    assert data["dependencies"][0]["target"] == "vpc-1"
    assert data["coverage"]["billed_without_collector"] == [
        {"cost_service": "Amazon Kinesis", "cost_90d": 3.0}
    ]

    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == ["Summary", "ec2", "s3", "Dependencies", "Errors"]
    ec2_rows = list(workbook["ec2"].iter_rows(values_only=True))
    header = ec2_rows[0]
    row = dict(zip(header, ec2_rows[1]))
    assert row["id"] == "i-1"
    assert row["is_active"] in (True, 1)
    assert row["dependencies"] == "in-vpc -> vpc-1"
    assert '"env": "prod"' in row["tags"]


class FakeAwsContext(Context):
    """Context whose real botocore clients are short-circuited with canned responses.

    Requests still go through botocore's parameter validation, so wrong operation names,
    paginators or parameters fail loudly; unlisted operations return an empty response.
    """

    def __init__(self, responses: dict[tuple[str, str], dict[str, Any]] | None = None) -> None:
        super().__init__(boto3.Session(), ACCOUNT)
        self.responses = responses or {}
        self._fake_clients: dict[tuple[str, str], Any] = {}

    def client(self, service: str, region: str | None = None) -> Any:
        key = (service, region or self.global_region)
        if key not in self._fake_clients:
            fake = boto3.client(
                service,
                region_name=key[1],
                aws_access_key_id="test",
                aws_secret_access_key="test",  # pragma: allowlist secret
            )

            def handler(model: Any, **_: Any) -> tuple[AWSResponse, dict[str, Any]]:
                canned = self.responses.get((service, model.name), {})
                return AWSResponse("http://fake", 200, {}, None), canned

            fake.meta.events.register("before-call", handler)
            self._fake_clients[key] = fake
        return self._fake_clients[key]


@pytest.mark.parametrize("service", sorted(COLLECTORS))
def test_every_collector_runs_cleanly_against_empty_account(service: str) -> None:
    ctx = FakeAwsContext()
    assert COLLECTORS[service].collect(ctx, REGION) == []
    assert ctx.errors == []


def test_ec2_collector_maps_instance_state_dependencies_and_activity() -> None:
    launched = datetime(2026, 1, 2, tzinfo=timezone.utc)
    seen = datetime(2026, 10, 26, tzinfo=timezone.utc)
    ctx = FakeAwsContext(
        {
            ("ec2", "DescribeInstances"): {
                "Reservations": [
                    {
                        "Instances": [
                            {
                                "InstanceId": "i-run",
                                "State": {"Name": "running"},
                                "LaunchTime": launched,
                                "InstanceType": "t3.micro",
                                "VpcId": "vpc-1",
                                "SubnetId": "subnet-1",
                                "SecurityGroups": [{"GroupId": "sg-1"}],
                                "Tags": [{"Key": "Name", "Value": "web"}],
                                "IamInstanceProfile": {"Arn": "arn:aws:iam::1:instance-profile/p"},
                                "BlockDeviceMappings": [{"Ebs": {"VolumeId": "vol-1"}}],
                            },
                            {"InstanceId": "i-off", "State": {"Name": "stopped"}},
                        ]
                    }
                ]
            },
            ("cloudwatch", "GetMetricData"): {
                "MetricDataResults": [{"Id": "q0", "Timestamps": [seen], "Values": [10.0]}]
            },
        }
    )
    resources = {r.id: r for r in COLLECTORS["ec2"].collect(ctx, REGION)}

    running = resources["i-run"]
    assert running.name == "web"
    assert running.is_active is True
    assert running.created_at == launched.isoformat()
    assert running.arn == f"arn:aws:ec2:{REGION}:{ACCOUNT}:instance/i-run"
    assert {d.relation for d in running.dependencies} >= {
        "in-vpc",
        "in-subnet",
        "security-group",
        "attached-volume",
        "instance-profile",
    }
    assert resources["i-off"].is_active is False
    assert ctx.errors == []


def test_lambda_collector_activity_based_active_flag() -> None:
    fn_arn = f"arn:aws:lambda:{REGION}:{ACCOUNT}:function:hello"
    queue_arn = f"arn:aws:sqs:{REGION}:{ACCOUNT}:jobs"
    ctx = FakeAwsContext(
        {
            ("lambda", "ListFunctions"): {
                "Functions": [
                    {
                        "FunctionName": "hello",
                        "FunctionArn": fn_arn,
                        "Runtime": "python3.12",
                        "Role": f"arn:aws:iam::{ACCOUNT}:role/r",
                        "Environment": {
                            "Variables": {"SECRET": "do-not-leak"},  # pragma: allowlist secret
                        },
                    }
                ]
            },
            ("lambda", "ListEventSourceMappings"): {
                "EventSourceMappings": [
                    {"FunctionArn": fn_arn, "EventSourceArn": queue_arn, "State": "Enabled"}
                ]
            },
            ("lambda", "ListTags"): {"Tags": {"team": "x"}},
            # No datapoints -> observed, but not active.
            ("cloudwatch", "GetMetricData"): {"MetricDataResults": [{"Id": "q0"}]},
        }
    )
    (function,) = COLLECTORS["lambda"].collect(ctx, REGION)
    assert function.is_active is False
    assert function.last_activity is None
    assert function.tags == {"team": "x"}
    assert function.metadata["environment_variable_count"] == 1
    assert "do-not-leak" not in str(function.metadata)
    assert Dependency("triggered-by", queue_arn, "sqs") in function.dependencies


def test_run_inventory_end_to_end_with_discovery_and_cost(tmp_path: Path) -> None:
    from src.aws_inventory.runner import run_inventory

    kinesis = f"arn:aws:kinesis:{REGION}:{ACCOUNT}:stream/events"
    ctx = FakeAwsContext(
        {
            ("resourcegroupstaggingapi", "GetResources"): {
                "ResourceTagMappingList": [
                    {"ResourceARN": kinesis, "Tags": [{"Key": "team", "Value": "data"}]}
                ]
            },
            ("ce", "GetCostAndUsage"): {
                "ResultsByTime": [
                    {
                        "Groups": [
                            {
                                "Keys": ["Amazon Kinesis", REGION],
                                "Metrics": {"UnblendedCost": {"Amount": "7.5", "Unit": "USD"}},
                            },
                            {
                                "Keys": ["AWS Lambda", REGION],
                                "Metrics": {"UnblendedCost": {"Amount": "1.25", "Unit": "USD"}},
                            },
                        ]
                    }
                ]
            },
            ("iam", "ListAccountAliases"): {"AccountAliases": ["demo"]},
        }
    )
    inventory = run_inventory(ctx, [REGION])

    assert inventory.account["alias"] == "demo"
    assert inventory.errors == []
    (stub,) = inventory.resources
    assert (stub.service, stub.status, stub.tags) == ("kinesis", "unknown", {"team": "data"})
    data = inventory.to_dict()
    assert data["cost"]["by_service"]["lambda"]["total"] == 1.25
    assert data["coverage"]["billed_without_collector"] == [
        {"cost_service": "Amazon Kinesis", "cost_90d": 7.5}
    ]
    write_xlsx(inventory, tmp_path / "e2e.xlsx")
    write_json(inventory, tmp_path / "e2e.json")
