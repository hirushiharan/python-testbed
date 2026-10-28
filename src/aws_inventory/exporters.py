"""Write the inventory as JSON (full fidelity) and Excel (flattened, one sheet per service)."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.aws_inventory.model import Inventory

EXCEL_CELL_LIMIT = 32000
MAX_COLUMN_WIDTH = 60
HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")

RESOURCE_COLUMNS = [
    "type",
    "id",
    "name",
    "region",
    "status",
    "is_active",
    "created_at",
    "last_activity",
    "last_activity_source",
    "arn",
    "dependency_count",
    "dependencies",
    "tags",
    "metadata",
]
SUMMARY_COLUMNS = [
    "service",
    "resource_count",
    "active_count",
    "inactive_count",
    "unknown_activity_count",
    "regions",
    "status_counts",
    "last_activity",
    "cost_90d",
]
DEPENDENCY_COLUMNS = [
    "source",
    "source_service",
    "source_type",
    "relation",
    "target",
    "target_service",
]
ERROR_COLUMNS = ["service", "region", "operation", "message"]


def write_json(inventory: Inventory, path: Path) -> Path:
    """Write the nested inventory to ``path``."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(inventory.to_dict(), indent=2, default=str), encoding="utf-8")
    return path


def write_xlsx(inventory: Inventory, path: Path) -> Path:
    """Write Summary, Dependencies, Errors and one sheet per service to ``path``."""
    path.parent.mkdir(parents=True, exist_ok=True)
    data = inventory.to_dict()
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    _fill(
        summary,
        SUMMARY_COLUMNS,
        [
            {
                "service": name,
                **{k: block["summary"].get(k) for k in SUMMARY_COLUMNS if k in block["summary"]},
                "regions": ", ".join(
                    f"{r} ({c})" for r, c in block["summary"]["by_region"].items()
                ),
                "status_counts": ", ".join(
                    f"{s}: {c}" for s, c in block["summary"]["status_counts"].items()
                ),
            }
            for name, block in data["services"].items()
        ],
    )

    used_names = {"Summary", "Dependencies", "Errors"}
    for name, block in data["services"].items():
        sheet = workbook.create_sheet(_sheet_name(name, used_names))
        _fill(sheet, RESOURCE_COLUMNS, [_flatten_resource(r) for r in block["resources"]])

    _fill(workbook.create_sheet("Dependencies"), DEPENDENCY_COLUMNS, data["dependencies"])
    _fill(workbook.create_sheet("Errors"), ERROR_COLUMNS, data["coverage"]["errors"])

    workbook.save(path)
    return path


def _flatten_resource(resource: dict[str, Any]) -> dict[str, Any]:
    flat = dict(resource)
    flat["dependency_count"] = len(resource["dependencies"])
    flat["dependencies"] = "; ".join(
        f"{d['relation']} -> {d['target']}" for d in resource["dependencies"]
    )
    flat["tags"] = json.dumps(resource["tags"], default=str) if resource["tags"] else ""
    flat["metadata"] = json.dumps(resource["metadata"], default=str) if resource["metadata"] else ""
    return flat


def _sheet_name(service: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", service)[:31] or "service"
    name, counter = base, 2
    while name in used:
        name = f"{base[:28]}_{counter}"
        counter += 1
    used.add(name)
    return name


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, (int, float)):
        return value
    text = value if isinstance(value, str) else json.dumps(value, default=str)
    return text[:EXCEL_CELL_LIMIT]


def _fill(sheet: Any, columns: list[str], rows: list[dict[str, Any]]) -> None:
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append([_cell(row.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    for index, column in enumerate(columns, start=1):
        longest = max([len(column)] + [len(str(row.get(column) or "")) for row in rows[:200]])
        sheet.column_dimensions[get_column_letter(index)].width = min(longest + 2, MAX_COLUMN_WIDTH)
